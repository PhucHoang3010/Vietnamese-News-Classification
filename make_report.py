"""Sinh báo cáo Word từ kết quả trong outputs/ và outputs/experiments/:  python make_report.py"""
import json
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

O = Path(__file__).parent / "outputs"
E = O / "experiments"
m = json.load(open(O / "metrics.json", encoding="utf-8")); bench = pd.read_csv(O / "so_sanh_mo_hinh.csv", encoding="utf-8")
rep = pd.read_csv(O / "bao_cao_phan_loai.csv", index_col=0, encoding="utf-8-sig"); kw = json.load(open(O / "tu_khoa_dac_trung.json", encoding="utf-8"))
err = pd.read_csv(O / "mau_du_doan_sai.csv", encoding="utf-8-sig")


def rcsv(name):
    return pd.read_csv(E / name, encoding="utf-8-sig")


def rjson(name):
    return json.load(open(E / name, encoding="utf-8"))


tun = rcsv("tuning_final.csv"); svmj = rjson("tuning_results_best_svm.json"); sig = rcsv("significance_vs_svm.csv")
stab = rcsv("stability_summary.csv"); sp = rcsv("stability_splits.csv"); ci = rjson("bootstrap_ci.json")["Linear SVM"]; nd = rjson("near_duplicates.json")
abl = rcsv("ablation_results.csv").set_index("bien_the"); lc = rcsv("learning_curve.csv"); cal = rcsv("calibration_results.csv"); rc = rcsv("risk_coverage.csv")
ec = rjson("error_confidence.json"); pairs = rcsv("error_pairs.csv"); bylen = rcsv("error_by_length.csv"); conf_err = rcsv("errors_confident.csv"); e88 = rcsv("errors_88.csv")
fa = rjson("explain_faithfulness.json"); fcsv = rcsv("explain_faithfulness.csv")
qr = rcsv("qa_retrieval.csv"); qc = rcsv("qa_confidence.csv"); qa = rjson("qa_abstain.json"); T_cal = rjson("calibration_T.json")["T"]

d = Document(); d.styles["Normal"].font.name = "Times New Roman"; d.styles["Normal"].font.size = Pt(13)


def vn(x, k=1):
    """Số kiểu Việt Nam: dấu phẩy thập phân."""
    return f"{x:.{k}f}".replace(".", ",")


def pc(x, k=1):
    return vn(100 * x, k) + "%"


def pv(p):
    return "< 0,001" if p < 0.001 else vn(p, 3)


def th(n):
    return f"{int(n):,}".replace(",", ".")


def H(t, l=1): d.add_heading(t, l)
def P(t): d.add_paragraph(t)
def B(items): [d.add_paragraph(i, style="List Bullet") for i in items]
def IMG(f, cap, w=14, folder=O):
    d.add_picture(str(folder / f), width=Cm(w)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def TABLE(df, fmt="{:.3f}"):
    t = d.add_table(rows=1, cols=len(df.columns)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns): t.rows[0].cells[i].text = str(c)
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(r): cells[i].text = fmt.format(v).replace(".", ",") if isinstance(v, float) else str(v)
    for row in t.rows:                                   # chữ trong bảng nhỏ hơn để ít bị xuống dòng
        for c in row.cells:
            for p_ in c.paragraphs:
                for run in p_.runs: run.font.size = Pt(10.5)
    d.add_paragraph()


t = d.add_heading("ĐỒ ÁN MÔN HỌC MÁY\nPHÂN LOẠI TIN TỨC TIẾNG VIỆT", 0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
P("Sinh viên thực hiện: ........................   MSSV: ..............\nLớp: ..............   Giảng viên: ........................")

# ------------------------------------------------------------------ 1
H("1. Giới thiệu")
P("Mỗi ngày các trang báo điện tử đăng hàng nghìn bài viết. Việc gán chủ đề (giáo dục, sức khỏe, thể thao, pháp luật...) thủ công tốn nhiều thời gian. "
  "Đề tài xây dựng hệ thống học máy tự động phân loại bài báo tiếng Việt vào 11 chủ đề dựa trên nội dung văn bản.")
P("Mục tiêu: (1) xây dựng pipeline tiền xử lý – trích xuất đặc trưng – huấn luyện; (2) so sánh công bằng một số thuật toán học máy cổ điển (mỗi mô hình đều được chỉnh tham số bằng cùng một quy trình); "
  "(3) đánh giá bằng Accuracy, Precision, Recall, F1 kèm khoảng tin cậy, kiểm định thống kê, độ ổn định khi chia lại dữ liệu và ablation; (4) giải thích được vì sao mô hình chọn chủ đề và biết độ tin cậy của dự đoán; "
  "(5) cung cấp chương trình chạy được trên máy cấu hình thấp (CPU, không cần GPU hay mô hình ngôn ngữ lớn) và đóng gói thành website tin tức.")

# ------------------------------------------------------------------ 2
H("2. Dữ liệu")
P(f"Bộ dữ liệu gồm {th(m['n_samples'])} bài báo tiếng Việt, {m['n_classes']} chủ đề, mỗi chủ đề 500 bài (cân bằng). "
  "Mỗi mẫu có tiêu đề, mô tả (sapo), nội dung và nhãn (chuyên mục của bài trên trang báo gốc). Văn bản đầu vào = tiêu đề + mô tả + nội dung. "
  "Dữ liệu chia phân tầng (stratified) theo tỉ lệ 70% huấn luyện (3.850) / 15% kiểm định (825) / 15% kiểm thử (825), random_state=42. Không có bài trùng hoàn toàn văn bản.")
P("Kiểm tra rò rỉ do bài gần giống nhau: với mỗi bài kiểm thử, chúng em tính độ giống cosine (TF-IDF) lớn nhất với các bài huấn luyện. "
  f"Trung vị {vn(nd['median'], 2)}, phân vị 90% là {vn(nd['p90'], 2)}, lớn nhất {vn(nd['max'], 2)}; chỉ {nd['n_ge_0.5']}/{nd['n_test']} bài có độ giống từ 0,5 trở lên (mô hình đoán đúng cả {nd['n_ge_0.5']} bài này) "
  f"và độ chính xác trên các bài còn lại là {pc(nd['acc_lt_0.5'])}. Vì vậy kết quả kiểm thử không bị thổi phồng bởi các bài gần trùng.")
IMG("phan_bo_nhan.png", "Hình 1. Phân bố số bài theo chủ đề", 12)
IMG("do_dai_bai.png", "Hình 2. Phân bố độ dài bài báo", 12)

# ------------------------------------------------------------------ 3
H("3. Phương pháp")
H("3.1 Tiền xử lý", 2)
B(["Chuẩn hóa Unicode NFC, chuyển chữ thường.", "Loại bỏ URL, email, chữ số và dấu câu; giữ nguyên dấu tiếng Việt.", "Gộp khoảng trắng; loại mẫu quá ngắn và mẫu trùng lặp."])
H("3.2 Trích xuất đặc trưng TF-IDF", 2)
P("Văn bản được biểu diễn bằng vector TF-IDF với unigram + bigram (bigram giúp bắt cụm từ tiếng Việt như “học sinh”, “ngân hàng”, “diễn viên”), "
  "min_df=2, max_df=0.9, sublinear_tf=True, tối đa 200.000 đặc trưng. TF-IDF cho trọng số cao với từ đặc trưng của một chủ đề và thấp với từ xuất hiện ở mọi chủ đề.")
H("3.3 Các mô hình so sánh", 2)
B(["Đường cơ sở: luôn đoán chủ đề phổ biến nhất (chỉ đúng 9,1% vì 11 chủ đề cân bằng).",
   "Complement Naive Bayes – nhanh, phù hợp văn bản.", "k-NN với độ đo cosine.", "Logistic Regression.",
   "Linear SVM (LinearSVC) – thường mạnh nhất với TF-IDF thưa, nhiều chiều.", "Random Forest.",
   "Mạng nơ-ron nhỏ: TF-IDF → giảm chiều SVD 300 → MLP 256 nơ-ron (chạy được trên CPU)."])
H("3.4 Quy trình đánh giá và chống rò rỉ", 2)
B(["Tập kiểm thử (825 bài) chỉ dùng để báo cáo kết quả cuối; mọi việc chọn mô hình và tham số chỉ dùng cross-validation 5-fold trên tập huấn luyện (3.850 bài).",
   "TF-IDF được học lại trong từng fold bằng Pipeline nên không rò rỉ thông tin từ fold kiểm tra.",
   "Ablation dùng 5-fold CV trên huấn luyện + kiểm định (4.675 bài), không dùng tập kiểm thử.",
   "Mọi thí nghiệm nằm trong thư mục experiments/ và chạy lại được; kết quả lưu ở outputs/experiments/."])

# ------------------------------------------------------------------ 4
H("4. Kết quả thực nghiệm")
H("4.1 So sánh mô hình công bằng", 2)
P("Lần so sánh nhanh đầu tiên (train.py) dùng tham số cố định cho cả bốn mô hình và chỉ tinh chỉnh riêng Linear SVM, nên có thể thiên vị mô hình được chọn. "
  "Vì vậy chúng em làm lại: mỗi mô hình đều được chỉnh siêu tham số bằng GridSearchCV 5-fold (Macro-F1) trên tập huấn luyện; khi tham số tốt nhất nằm ở biên lưới thì mở rộng lưới (Logistic Regression: C đến 300; Linear SVM: C từ 0,1 đến 64). "
  "Tập kiểm định và kiểm thử không tham gia chọn tham số.")
tt = tun[~tun.model.str.startswith("Đoán")].sort_values("cv_macro_f1_mean", ascending=False)
tb = pd.DataFrame({"Mô hình": tt.model.str.replace(r" \(TF-IDF.*", "", regex=True), "Macro-F1 CV (5-fold)": [f"{vn(a, 3)} ± {vn(b, 3)}" for a, b in zip(tt.cv_macro_f1_mean, tt.cv_macro_f1_std)],
                   "Macro-F1 kiểm định": [vn(x, 3) for x in tt.val_macro_f1], "Accuracy kiểm thử": [vn(x, 3) for x in tt.test_acc],
                   "Huấn luyện (s)": [vn(x, 0) for x in tt.fit_sec], "Kích thước (MB)": [vn(x, 1) for x in tt.model_mb]})
TABLE(tb)
base_row = tun[tun.model.str.startswith("Đoán")].iloc[0]
P(f"Đường cơ sở (đoán chủ đề phổ biến nhất) chỉ đạt Accuracy {pc(base_row.test_acc)}. Linear SVM có Macro-F1 CV cao nhất ({vn(float(tt.iloc[0].cv_macro_f1_mean), 4)}) nhưng chỉ nhỉnh Logistic Regression ({vn(float(tt.iloc[1].cv_macro_f1_mean), 4)}); "
  f"so với lần so sánh nhanh ban đầu (SVM hơn Logistic Regression khoảng 1,3 điểm phần trăm trên tập kiểm định) thì khoảng cách gần như biến mất khi Logistic Regression được chỉnh C công bằng. "
  f"Với Linear SVM, Macro-F1 CV gần như phẳng khi C từ 1 đến 64 ({vn(svmj['plateau_cv_min'], 4)} – {vn(svmj['plateau_cv_max'], 4)}); C tốt nhất là {svmj['best_C']} nhưng chỉ hơn C = 2 khoảng {vn((svmj['best_cv'] - svmj['chosen_cv']) * 100, 2)} điểm, "
  "nhỏ hơn độ lệch chuẩn giữa các fold (0,41 điểm) nên giữ C = 2 (đơn giản, chính quy hóa mạnh hơn).")
IMG("tuning_results.png", "Hình 3. Kết quả chỉnh tham số 5-fold CV của các mô hình", 14, E)
H("Kiểm định thống kê giữa các mô hình", 3)
P("Trên cùng 825 bài kiểm thử, chúng em so sánh từng cặp bằng kiểm định McNemar (xét các bài mà hai mô hình trả lời khác nhau) và khoảng tin cậy 95% bootstrap (2.000 lần lấy mẫu lại) cho chênh lệch Accuracy.")
sg = pd.DataFrame({"So Linear SVM với": sig.vs.str.replace(r" \(TF-IDF.*", "", regex=True), "Chỉ SVM đúng": sig.svm_only_right.astype(str), "Chỉ mô hình kia đúng": sig.other_only_right.astype(str),
                   "p (McNemar)": [pv(x) for x in sig.mcnemar_p], "Chênh Accuracy, điểm % [95% CI]": [f"{vn(a * 100, 1)} [{vn(b * 100, 1)}; {vn(c * 100, 1)}]" for a, b, c in zip(sig.acc_diff, sig.diff_lo, sig.diff_hi)]})
TABLE(sg)
P("Kết luận: Linear SVM tốt hơn có ý nghĩa thống kê so với k-NN và Random Forest (p < 0,05), nhưng không tách khỏi Logistic Regression, Complement Naive Bayes và MLP nhỏ (p > 0,1). "
  "Chúng em chọn Linear SVM vì có điểm trung bình cao nhất, huấn luyện nhanh hơn Logistic Regression khoảng 4 lần, file mô hình nhỏ (15,9 MB) và là mô hình tuyến tính nên giải thích được trực tiếp (mục 6.4), "
  "chứ không khẳng định nó vượt trội hơn Logistic Regression.")

H("4.2 Kết quả cuối trên tập kiểm thử", 2)
P(f"Trên {m['test_size']} mẫu kiểm thử (chưa từng dùng khi huấn luyện/chọn mô hình): Accuracy = {vn(m['test_accuracy'], 4)} ({int(round(m['test_accuracy'] * m['test_size']))}/{m['test_size']} bài), Macro-F1 = {vn(m['test_macro_f1'], 4)}. "
  f"Khoảng tin cậy 95% (bootstrap) cho Accuracy là [{pc(ci['acc_lo'])}; {pc(ci['acc_hi'])}] và cho Macro-F1 là [{vn(ci['f1_lo'], 3)}; {vn(ci['f1_hi'], 3)}]. "
  "Ba con số Macro-F1 gần nhau (CV 5-fold 0,884; kiểm định 0,886; kiểm thử 0,893), không thấy dấu hiệu chọn mô hình quá khớp vào tập kiểm định. "
  "Lưu ý: độ chính xác trên tập huấn luyện là 100% (SVM trong không gian 200.000 chiều tách hết dữ liệu học), nên điều đáng tin là kết quả trên dữ liệu chưa thấy.")
r = rep.loc[[i for i in rep.index if i not in ("accuracy", "macro avg", "weighted avg")], ["precision", "recall", "f1-score", "support"]].reset_index()
r.columns = ["Chủ đề", "Precision", "Recall", "F1", "Số mẫu"]; r["Số mẫu"] = r["Số mẫu"].astype(int).astype(str); TABLE(r, "{:.3f}")
IMG("ma_tran_nham_lan.png", "Hình 4. Ma trận nhầm lẫn trên tập kiểm thử", 15)

H("4.3 Độ ổn định khi chia lại dữ liệu", 2)
P("Kết quả trên một lần chia dữ liệu có thể may rủi. Chúng em lặp lại phép chia 70/15/15 với 10 hạt giống ngẫu nhiên khác nhau (seed 0–9), huấn luyện lại và đo trên tập kiểm thử của mỗi lần chia; "
  "ngoài ra chạy 5-fold CV lặp 3 lần trên toàn bộ 5.500 bài.")
sb = pd.DataFrame({"Mô hình": stab.model, "Accuracy (10 lần chia)": [f"{vn(a * 100, 2)} ± {vn(b * 100, 2)}" for a, b in zip(stab.split_acc_mean, stab.split_acc_std)],
                   "Macro-F1 (10 lần chia)": [f"{vn(a, 4)} ± {vn(b, 4)}" for a, b in zip(stab.split_f1_mean, stab.split_f1_std)],
                   "Nhỏ nhất – lớn nhất": [f"{vn(a, 3)} – {vn(b, 3)}" for a, b in zip(stab.split_f1_min, stab.split_f1_max)],
                   "Macro-F1 (CV 5-fold × 3)": [f"{vn(a, 4)} ± {vn(b, 4)}" for a, b in zip(stab.cv_f1_mean, stab.cv_f1_std)]})
TABLE(sb)
sv_ = stab[stab.model == "Linear SVM"].iloc[0]
P(f"Linear SVM đạt Accuracy {vn(sv_.split_acc_mean * 100, 2)}% ± {vn(sv_.split_acc_std * 100, 2)} trên 10 lần chia, kết quả chính thức ({pc(m['test_accuracy'], 2)}) nằm trong dải này. "
  f"Trong {int((sp[sp.model == 'Linear SVM'].set_index('seed').macro_f1 > sp[sp.model == 'Logistic Regression'].set_index('seed').macro_f1).sum())}/10 lần chia Linear SVM có Macro-F1 cao hơn Logistic Regression, nhưng chênh lệch nhỏ (cùng cỡ độ lệch chuẩn). "
  "Complement Naive Bayes dao động nhiều hơn hẳn (độ lệch chuẩn gấp đôi).")
IMG("stability.png", "Hình 5. Macro-F1 trên tập kiểm thử của 10 phép chia ngẫu nhiên", 13, E)

H("4.4 Ablation: từng thành phần đóng góp bao nhiêu?", 2)
P("Mỗi lần chỉ đổi một thành phần so với cấu hình chuẩn, đo bằng 5-fold CV trên huấn luyện + kiểm định (không dùng tập kiểm thử); p là kiểm định t ghép cặp trên 5 fold (ít fold nên kiểm định có độ mạnh thấp: “chưa có ý nghĩa” không có nghĩa là “không có tác dụng”).")
sel = [("Cấu hình chuẩn", "Cấu hình chuẩn (TF-IDF 1–2 gram, log-tf, IDF, tiêu đề + sapo + nội dung)"), ("Chỉ tiêu đề", "Chỉ dùng tiêu đề"), ("Tiêu đề + sapo", "Tiêu đề + sapo"), ("Chỉ nội dung", "Chỉ dùng nội dung"),
       ("Chỉ từ đơn (1-gram)", "Chỉ từ đơn (bỏ cặp từ)"), ("1–3 gram", "Thêm bộ ba từ (1–3 gram)"), ("Bỏ log-tf (sublinear_tf=False)", "Bỏ log-tf"), ("Bỏ IDF (chỉ TF)", "Bỏ IDF"),
       ("Tối đa 5.000 đặc trưng", "Chỉ giữ 5.000 từ"), ("Tối đa 20.000 đặc trưng", "Chỉ giữ 20.000 từ"), ("min_df = 5", "Loại từ hiếm hơn (min_df = 5)"), ("Bỏ dấu tiếng Việt", "Bỏ dấu tiếng Việt"),
       ("Tách từ (pyvi) + 1–2 gram", "Tách từ pyvi + 1–2 gram")]
ab = pd.DataFrame({"Biến thể": [b for _, b in sel], "Macro-F1 (%)": [f"{vn(abl.loc[a, 'macro_f1_mean'] * 100, 2)} ± {vn(abl.loc[a, 'macro_f1_std'] * 100, 2)}" for a, _ in sel],
                   "Chênh so với chuẩn (điểm %)": ["—" if a == "Cấu hình chuẩn" else f"{abl.loc[a, 'delta_vs_base'] * 100:+.2f}".replace(".", ",") for a, _ in sel],
                   "p (t ghép cặp)": ["—" if a == "Cấu hình chuẩn" else pv(abl.loc[a, "paired_t_p"]) for a, _ in sel]})
TABLE(ab)
IMG("ablation.png", "Hình 6. Ablation (5-fold CV, không dùng tập kiểm thử)", 15, E)
B([f"Nội dung bài báo là phần quan trọng nhất: chỉ dùng tiêu đề chỉ đạt Macro-F1 {vn(abl.loc['Chỉ tiêu đề', 'macro_f1_mean'] * 100, 1)}%, tiêu đề + sapo đạt {vn(abl.loc['Tiêu đề + sapo', 'macro_f1_mean'] * 100, 1)}%, "
   f"trong khi chỉ dùng nội dung ({vn(abl.loc['Chỉ nội dung', 'macro_f1_mean'] * 100, 1)}%) đã ngang cấu hình đầy đủ ({vn(abl.loc['Cấu hình chuẩn', 'macro_f1_mean'] * 100, 1)}%).",
   "Cặp từ (bigram), log-tf và vốn từ đủ lớn đều đóng góp có ý nghĩa thống kê (bỏ đi làm Macro-F1 giảm 0,76 – 1,03 điểm, p < 0,05); từ 20.000 đặc trưng trở lên không còn khác biệt.",
   "Bỏ IDF giảm 0,49 điểm nhưng chưa có ý nghĩa (p = 0,10); bỏ dấu tiếng Việt giảm 0,22 điểm (p = 0,27).",
   "Tách từ bằng pyvi và thêm bộ ba từ chỉ tăng 0,13 – 0,24 điểm, chưa có ý nghĩa (p = 0,24 và 0,33) nên không đưa vào mô hình chính."])

H("4.5 Đường cong học", 2)
P("Huấn luyện Linear SVM trên các phần ngày càng lớn của tập huấn luyện (mỗi mốc lấy mẫu con ngẫu nhiên 5 lần), đo Macro-F1 trên tập kiểm định.")
lt = pd.DataFrame({"Số bài huấn luyện": [f"{int(n):,}".replace(",", ".") for n in lc.n], "Accuracy huấn luyện": [pc(x, 0) for x in lc.train_acc], "Macro-F1 kiểm định": [pc(x, 1) for x in lc.val_f1]})
TABLE(lt)
IMG("learning_curve.png", "Hình 7. Đường cong học", 12, E)
P(f"Macro-F1 tăng từ {pc(lc.val_f1.iloc[0])} (192 bài) lên {pc(lc.val_f1.iloc[-1])} (3.850 bài) và chưa phẳng hẳn: từ 2.887 lên 3.850 bài vẫn tăng khoảng {vn((lc.val_f1.iloc[-1] - lc.val_f1.iloc[-2]) * 100, 1)} điểm. "
  "Độ chính xác trên tập huấn luyện luôn 100% nên khoảng cách huấn luyện – kiểm định là phương sai (thiếu dữ liệu), gợi ý rằng thêm dữ liệu sẽ còn giúp nhưng lợi ích giảm dần.")

H("4.6 Hiệu chỉnh độ tin cậy và dự đoán có chọn lọc", 2)
P("Website hiển thị “mức nghiêng” (%) của AI cho từng chuyên mục, tính bằng softmax(T × điểm SVM). Một độ tin cậy hữu ích khi 90% nghĩa là khoảng 90% những lần như vậy đúng. "
  "Chúng em đo sai lệch hiệu chỉnh ECE (chia 10 nhóm độ tin cậy), Brier và NLL trên tập kiểm thử; hệ số T được chọn trên tập kiểm định (temperature scaling).")
ct = pd.DataFrame({"Cách tính": cal.method, "Độ tin cậy trung bình": [pc(x, 1) for x in cal.mean_conf], "ECE": [pc(x, 1) for x in cal.ECE], "Brier": [vn(x, 3) for x in cal.brier], "NLL": [vn(x, 3) for x in cal.nll]})
TABLE(ct)
c0, c1 = cal.iloc[0], cal.iloc[1]
P(f"Hệ số ban đầu (T = 3) làm AI quá khiêm tốn: độ tin cậy trung bình {pc(c0.mean_conf)} trong khi độ chính xác thật là {pc(c0.acc)} (ECE {pc(c0.ECE)}). "
  f"Chọn T = {vn(T_cal, 1)} trên tập kiểm định đưa ECE xuống {pc(c1.ECE)} và NLL từ {vn(c0.nll, 3)} xuống {vn(c1.nll, 3)}, tốt hơn Platt scaling (ECE {pc(cal.iloc[2].ECE)}) và gần bằng Logistic Regression ({pc(cal.iloc[3].ECE)}). "
  f"Website dùng T = {vn(T_cal, 1)}. Các bài AI đoán đúng có mức nghiêng trung bình {vn(ec['pct_correct'], 1)}%, còn bài đoán sai chỉ {vn(ec['pct_wrong'], 1)}%.")
IMG("calibration.png", "Hình 8. Biểu đồ độ tin cậy (chỉ vẽ nhóm có từ 10 bài; kích thước chấm tỉ lệ số bài)", 15, E)
P("Dự đoán có chọn lọc: nếu AI chỉ tự quyết những bài nó chắc chắn nhất và chuyển phần còn lại cho biên tập viên thì độ chính xác trên phần tự quyết tăng như sau (xếp theo chênh lệch giữa hai chuyên mục có xác suất cao nhất).")
rt = pd.DataFrame({"Tỉ lệ bài AI tự quyết": [pc(x, 0) for x in rc.coverage], "Số bài": rc.n.astype(str), "Accuracy trên phần tự quyết": [pc(x, 1) for x in rc.accuracy]})
TABLE(rt)

# ------------------------------------------------------------------ 5
H("5. Phân tích lỗi")
n_ts = int(((e88.truth == "Thời sự") | (e88.pred == "Thời sự")).sum())
P(f"Mô hình sai {len(err)}/{m['test_size']} bài ({pc(len(err) / m['test_size'])}). Chủ đề đặc trưng rõ như Thể thao, Sức khỏe, Xe, Giải trí, Pháp luật đạt F1 > 0,92; “Thời sự”, “Đời sống”, “Kinh doanh”, “Thế giới” thấp hơn. "
  f"Chuyên mục Thời sự có mặt trong {n_ts}/{len(err)} bài sai (là nhãn thật hoặc nhãn dự đoán). Các cặp nhầm nhiều nhất:")
pp = pairs.head(6).copy(); pp.columns = ["Nhãn thật", "AI dự đoán", "Số bài"]; pp["Số bài"] = pp["Số bài"].astype(str); TABLE(pp)
P(f"Trong {ec['n_wrong']} bài sai, chuyên mục đúng đứng thứ hai trong bảng xếp hạng của AI ở {ec['truth_rank2']} bài ({pc(ec['truth_rank2'] / ec['n_wrong'], 0)}) và từ thứ ba trở đi ở {ec['truth_rank3plus']} bài: phần lớn lỗi là “suýt đúng”. "
  f"Không thấy xu hướng rõ theo độ dài bài: Accuracy của bốn nhóm độ dài (từ ngắn đến dài, mỗi nhóm khoảng 205 bài) lần lượt là " + ", ".join(pc(x, 1) for x in bylen.acc) + ".")
P(f"Có {ec['n_wrong_pct_ge_90']} bài AI sai nhưng rất tự tin (mức nghiêng từ 90% trở lên), trong tổng số {ec['n_all_pct_ge_90']} bài AI nghiêng từ 90% trở lên. Danh sách đầy đủ (để đọc và kiểm tra) trong outputs/experiments/errors_confident.csv; một số bài tiêu biểu:")
ce = conf_err.head(8).copy(); ce = pd.DataFrame({"Nhãn thật → AI": ce.truth + " → " + ce.pred, "Mức nghiêng": [vn(x, 1) + "%" for x in ce.pct], "Tiêu đề": ce.title})
TABLE(ce)
P("Đọc tiêu đề của các bài này cho thấy nhiều bài có thể xếp hợp lý vào cả chuyên mục gốc lẫn chuyên mục AI chọn: tin vụ án đăng ở mục Thời sự nhưng nội dung là Pháp luật, tin sinh viên giành huy chương thể thao đăng ở mục Giáo dục, tin tài chính ngân hàng có yếu tố pháp lý. "
  "Như vậy một phần “lỗi” là do ranh giới nhãn mờ chứ không hoàn toàn do mô hình, nhưng cũng có lỗi thật (ví dụ bản tin dự báo thời tiết bị xếp vào Kinh doanh với 91,4%). "
  "Chúng em chưa gán lại nhãn cho từng bài nên chưa thể định lượng tỉ lệ nhãn sai.")
P("Từ khóa đặc trưng nhất của từng chủ đề (trọng số lớn nhất của Linear SVM), giúp giải thích mô hình:")
TABLE(pd.DataFrame({"Chủ đề": list(kw), "Top từ khóa": [", ".join(v[:6]) for v in kw.values()]}))
P("Danh sách đầy đủ các bài dự đoán sai: outputs/mau_du_doan_sai.csv và outputs/experiments/errors_88.csv (kèm từ khóa ủng hộ nhãn thật và nhãn dự đoán).")

# ------------------------------------------------------------------ 6
H("6. Hệ thống web và các chức năng mở rộng")
P("Ngoài bài toán chính (phân loại), mô hình được đóng gói thành website tin tức VnNews (Flask) với các chức năng dùng những kỹ thuật học máy / xử lý ngôn ngữ khác nhau. "
  "Toàn bộ 5.500 bài được mô hình tự động gán chuyên mục; trang chủ có ô phân loại trực tiếp và mỗi bài có bảng “AI xếp bài này vào…” kèm mức nghiêng, top 3 chuyên mục và các từ khóa quyết định. "
  "Chỉ số độ chính xác chỉ tính trên tập kiểm thử; với bài thuộc tập huấn luyện/kiểm định, website ghi rõ AI đã được học bài đó.")
ev = json.load(open(O / "summary_eval.json", encoding="utf-8")) if (O / "summary_eval.json").exists() else {}
cl = json.load(open(O.parent / "data" / "clusters.json", encoding="utf-8")) if (O.parent / "data" / "clusters.json").exists() else {}
H("6.1 Tóm tắt tự động (TextRank)", 2)
P("Bài báo được tách thành câu (loại bỏ chú thích ảnh, dòng nguồn), mỗi câu được biểu diễn bằng vector TF-IDF. Đồ thị câu có cạnh là độ tương đồng cosine; chạy PageRank để tìm câu quan trọng nhất, "
  "cộng thêm điểm cho câu đầu bài và câu gần tiêu đề (Mihalcea & Tarau, 2004). Đây là tóm tắt trích rút nên không bịa thêm nội dung.")
if ev:
    t = d.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
    for i, c in enumerate(["Phương pháp", "ROUGE-1 F1", "ROUGE-2 F1", "Số bài"]): t.rows[0].cells[i].text = c
    for k, v in ev.items():
        cells = t.add_row().cells
        for i, x in enumerate([k, vn(v['rouge1'], 3), vn(v['rouge2'], 3), str(v['n'])]): cells[i].text = x
    d.add_paragraph()
    P("Đánh giá trên khoảng 400 bài, lấy phần sapo do biên tập viên viết làm bản tham chiếu (2 câu). TextRank cao hơn hai phương pháp cơ sở là Lead-2 (hai câu đầu) và chọn ngẫu nhiên, nhưng chênh lệch so với Lead-2 nhỏ.")
H("6.2 Gom cụm sự kiện / chủ đề con (học không giám sát)", 2)
sil = sum(c["sil"] for c in cl.values()) / len(cl) if cl else 0
P("Với mỗi chuyên mục, tiêu đề và sapo được biểu diễn bằng TF-IDF, giảm chiều bằng TruncatedSVD (LSA, 60 chiều) rồi phân cụm KMeans (k bằng số bài chia 26, khoảng 19 cụm mỗi chuyên mục). "
  "Tên cụm là các cụm từ đặc trưng nhất của cụm so với phần còn lại của chuyên mục; bài tiêu biểu là bài gần tâm cụm nhất. Phương pháp không dùng nhãn nên bổ sung cho bài toán phân loại có giám sát. "
  f"Độ gắn kết trung bình (silhouette) là {vn(sil, 3)}; giá trị thấp là bình thường với dữ liệu văn bản vì từ vựng giữa các câu chuyện giao thoa nhiều, nhưng nhãn cụm vẫn dễ hiểu (ví dụ “ung thư”, “bóng chuyền”, “ngân hàng”).")
H("6.3 Gợi ý bài đọc và bản tin nhanh", 2)
B(["Gợi ý “Dành cho bạn”: hệ gợi ý dựa trên nội dung (content-based). Vector sở thích là trung bình TF-IDF của các bài người dùng đã đọc (lưu trên trình duyệt); bài được gợi ý là các bài có cosine cao nhất chưa đọc.",
   "Bản tin nhanh 60 giây: tóm tắt hàng loạt, mỗi chuyên mục một tin nổi bật còn 2 câu, hiển thị mức rút gọn nội dung.",
   "Tìm kiếm: xếp hạng bài theo cosine TF-IDF giữa truy vấn và tiêu đề + sapo."])
H("6.4 AI Lab – giải thích dự đoán (Explainable AI)", 2)
P("Vì Linear SVM là mô hình tuyến tính nên điểm của mỗi chủ đề bằng tổng (trọng số × giá trị TF-IDF); có thể tách chính xác đóng góp của từng từ mà không cần xấp xỉ như LIME/SHAP: "
  "đóng góp(từ) = (w của chủ đề dự đoán − w trung bình các chủ đề khác) × TF-IDF(từ), cụm hai từ chia đều cho hai từ. "
  "Trang AI Lab tô đậm từng từ theo mức ủng hộ chủ đề dự đoán, cho phép bỏ từ để thử What-if, và phát hiện bài đa chủ đề bằng cách phân loại riêng từng đoạn.")
P("Kiểm chứng độ trung thực của lời giải thích (experiments/exp7_explain.py): với cả 825 bài kiểm thử, chúng em (a) xóa m từ mà AI cho là đóng góp nhiều nhất cho chủ đề dự đoán rồi dự đoán lại, "
  "(b) chỉ giữ lại m từ đó. Đối chứng gồm: m từ ngẫu nhiên theo loại từ, m từ ngẫu nhiên theo vị trí (từ hay gặp dễ trúng hơn) và m từ có TF-IDF lớn nhất. Đây là hai phép đo “comprehensiveness” và “sufficiency” thường dùng khi đánh giá giải thích (DeYoung và cs., 2020).")


def fcell(mode, mm, name):
    x = fa[f"{mode}_{mm}_{name}"]; return f"{vn(x['rate'] * 100, 1)}%"


names = ["AI chọn (đóng góp lớn nhất)", "TF-IDF lớn nhất", "Ngẫu nhiên (theo vị trí)", "Ngẫu nhiên (theo loại từ)"]
rows_ = []
for mm in (3, 5, 10):
    rows_.append([f"Xóa {mm} từ → bài đổi chuyên mục"] + [fcell("xoa", mm, n) for n in names])
for mm in (3, 5, 10):
    rows_.append([f"Chỉ giữ {mm} từ → bài GIỮ NGUYÊN chuyên mục"] + [f"{vn((1 - fa[f'giu_{mm}_{n}']['rate']) * 100, 1)}%" for n in names])
TABLE(pd.DataFrame(rows_, columns=["Phép thử"] + names))
x5 = fa["xoa_5_AI chọn (đóng góp lớn nhất)"]; g5 = fa["giu_5_AI chọn (đóng góp lớn nhất)"]
P(f"Kết quả: chỉ giữ 5 từ mà AI chỉ ra (khoảng {vn(g5['token_share'] * 100, 0)}% số từ của bài) thì {vn((1 - g5['rate']) * 100, 1)}% bài vẫn được xếp vào đúng chuyên mục cũ, trong khi 5 từ ngẫu nhiên chỉ giữ được {vn((1 - fa['giu_5_Ngẫu nhiên (theo loại từ)']['rate']) * 100, 0)}–{vn((1 - fa['giu_5_Ngẫu nhiên (theo vị trí)']['rate']) * 100, 0)}% và 5 từ có TF-IDF lớn nhất giữ được {vn((1 - fa['giu_5_TF-IDF lớn nhất']['rate']) * 100, 0)}%. "
  f"Ngược lại, xóa đúng 5 từ đó làm {vn(x5['rate'] * 100, 1)}% bài đổi chuyên mục (khoảng tin cậy 95%: {vn(x5['lo'] * 100, 1)}–{vn(x5['hi'] * 100, 1)}%), so với {fcell('xoa', 5, 'Ngẫu nhiên (theo loại từ)')}–{fcell('xoa', 5, 'TF-IDF lớn nhất')} khi xóa 5 từ chọn theo cách khác. "
  "Như vậy các từ được tô đậm thật sự mang quyết định của mô hình. Tuy nhiên quyết định dựa vào nhiều từ chứ không vài từ “then chốt”: phải xóa 20 từ mới làm gần một nửa số bài đổi chuyên mục "
  f"({vn(fa['ms'] and fcsv[(fcsv.che_do == 'xoa') & (fcsv.m == 20) & (fcsv.phuong_phap == names[0])].ty_le_doi_du_doan.iloc[0] * 100, 1)}%). Đây là kết quả trên 825 bài kiểm thử, không phải một ví dụ minh họa.")
IMG("explain_faithfulness.png", "Hình 9. Độ trung thực của lời giải thích (825 bài kiểm thử)", 15, E)
H("6.5 Hỏi AI về tin tức (hỏi đáp có dẫn nguồn)", 2)
P("Chức năng cho người đọc gõ một câu hỏi và nhận câu trả lời trích từ các bài báo, luôn kèm bài gốc và chuyên mục do AI phân loại; ô “Hỏi AI về bài này” trên từng bài chỉ tìm trong bài đó. "
  "Hệ thống không dùng mô hình sinh văn bản nên không tự bịa nội dung; hạn chế là chỉ trả lời được điều có sẵn trong kho 5.500 bài. Cách làm:")
B(["Truy hồi bài: điểm bài = 0,7 × cosine TF-IDF trên toàn văn (tiêu đề + sapo + nội dung, dùng chính TF-IDF của bộ phân loại) + 0,3 × cosine TF-IDF trên tiêu đề + sapo (chỉ mục của chức năng tìm kiếm); lấy 8 bài cao nhất.",
   "Chọn câu: tách các bài thành câu, xếp hạng câu theo cosine với câu hỏi (có thưởng nhẹ cho câu đầu bài), trả về tối đa 3 câu, mỗi bài không quá 2 câu.",
   "Độ chắc chắn dựa trên điểm của bài cao nhất: từ 0,30 là cao, từ 0,20 là trung bình, từ 0,10 là thấp; dưới 0,10 hệ thống trả lời “chưa tìm thấy thông tin phù hợp” thay vì đoán.",
   "Chủ đề của câu hỏi được suy ra bằng cách bỏ phiếu theo chuyên mục AI đã xếp cho các bài tìm được (phân loại trực tiếp câu hỏi ngắn không đủ tin cậy)."])
P("Đánh giá (experiments/exp8_qa.py): phép thử “tìm đúng bài” — dùng tiêu đề của mỗi bài làm câu hỏi, chỉ mục không chứa chính tiêu đề đó (tránh rò rỉ), xem bài đúng đứng thứ mấy trong kết quả trên 5.500 truy vấn.")
want = ["Chỉ sapo (TF-IDF tìm kiếm)", "Chỉ nội dung (TF-IDF của bộ phân loại)", "Kết hợp: nội dung 0.5 + sapo 0.5", "Kết hợp: nội dung 0.7 + sapo 0.3", "+ định tuyến chuyên mục (λ=0.5)"]
qq = qr.set_index("bien_the").loc[want].reset_index()
TABLE(pd.DataFrame({"Chỉ mục / cách chấm": [w.replace("0.", "0,").replace("λ=0,5", "λ = 0,5") for w in qq.bien_the], "Đúng bài ở vị trí 1": [pc(x, 1) for x in qq.recall_at_1], "Trong 5 kết quả đầu": [pc(x, 1) for x in qq.recall_at_5],
                    "Trong 10 kết quả đầu": [pc(x, 1) for x in qq.recall_at_10], "MRR": [vn(x, 3) for x in qq.mrr]}))
b7 = qr.set_index("bien_the").loc["Kết hợp: nội dung 0.7 + sapo 0.3"]
P(f"Kết hợp toàn văn và sapo với trọng số 0,7/0,3 cho bài đúng ở vị trí 1 trong {pc(b7.recall_at_1)} trường hợp và trong 5 kết quả đầu trong {pc(b7.recall_at_5)} (MRR {vn(b7.mrr, 3)}), cao hơn dùng riêng sapo hoặc riêng toàn văn. "
  "Định tuyến theo chuyên mục (chỉ ưu tiên bài thuộc chuyên mục mà bộ phân loại dự đoán cho câu hỏi) không cải thiện (vị trí 1: 86,3–86,6% so với 86,5% khi không định tuyến), nên không dùng; đây là một ví dụ cho thấy bộ phân loại không phải lúc nào cũng giúp các chức năng khác.")
P("Điểm của bài cao nhất có phản ánh độ đúng: tỉ lệ tìm đúng bài tăng đều theo khoảng điểm, nhờ đó đặt được nhãn “độ chắc chắn”:")
TABLE(pd.DataFrame({"Khoảng điểm của bài cao nhất": [k.replace(".", ",") for k in qc.khoang], "Số câu hỏi": [th(n) for n in qc.n], "Tìm đúng bài": [pc(x, 1) for x in qc.ty_le_dung_bai]}))
th10 = [x for x in qa["thresholds"] if abs(x["nguong"] - 0.10) < 1e-9][0]
P(f"Khả năng từ chối: đối chiếu 5.500 truy vấn thật với 5.500 truy vấn ngẫu nhiên (5 từ bất kỳ trong vốn từ), điểm của bài cao nhất tách hai nhóm với AUC = {vn(qa['auc_thuc_vs_ngau_nhien'], 3)}. "
  f"Với ngưỡng 0,10 đang dùng, hệ thống từ chối nhầm {pc(th10['tu_choi_nham_truy_van_that'])} truy vấn thật và vẫn trả lời nhầm {pc(th10['tra_loi_nham_truy_van_ngau_nhien'])} truy vấn ngẫu nhiên. "
  "Lưu ý: đây là phép thử tìm lại bài bằng tiêu đề, chưa phải đánh giá chất lượng câu trả lời với câu hỏi thật của người dùng; câu được trích có thể chưa đúng ý, vì vậy giao diện luôn dẫn nguồn để người đọc kiểm tra.")
H("6.6 Phân loại thời gian thực và chạy cục bộ", 2)
B(["Phân loại khi gõ: kết quả cập nhật theo thời gian thực sau mỗi lần ngừng gõ, không cần nhấn nút; trang chủ có ô dán văn bản và xem ngay chuyên mục AI chọn.",
   "Toàn bộ chức năng chạy cục bộ, không cần Internet hay thiết bị ngoại vi, nên ổn định khi trình bày; thêm tùy chọn python app.py --lan để máy khác cùng WiFi mở được."])

# ------------------------------------------------------------------ 7
H("7. Hạn chế và hướng phát triển")
B([f"Tách từ tiếng Việt: đã thử pyvi, Macro-F1 chỉ tăng {vn(abl.loc['Tách từ (pyvi) + 1–2 gram', 'delta_vs_base'] * 100, 2)} điểm (p = {vn(abl.loc['Tách từ (pyvi) + 1–2 gram', 'paired_t_p'], 2)}, chưa có ý nghĩa) nên không dùng; chưa thử underthesea/VnCoreNLP.",
   "Chưa thử mô hình ngôn ngữ lớn/PhoBERT vì mục tiêu chạy trên CPU; k-NN và MLP nhỏ đã thử nhưng không vượt Linear SVM.",
   "Dữ liệu chỉ 5.500 bài: đường cong học chưa phẳng nên thêm dữ liệu sẽ giúp; ranh giới nhãn mờ (mục 5) khiến khó vượt xa mức hiện tại nếu không làm sạch nhãn.",
   f"Phần trăm hiển thị đã hiệu chỉnh nhiệt độ (ECE {pc(c1.ECE)}) nhưng chưa phải xác suất tuyệt đối; chưa kiểm tra trên tin từ nguồn báo khác (ngoài miền dữ liệu).",
   "Hỏi AI chỉ trích câu có sẵn, không tổng hợp nhiều nguồn và chưa được đánh giá bằng câu hỏi của người dùng thật; chỉ trả lời được nội dung trong kho 5.500 bài (tĩnh, không cập nhật tin mới).",
   "Hướng phát triển: thu tin mới qua RSS để kiểm tra ngoài dữ liệu hiện có; gán lại nhãn cho các bài nghi sai; thử PhoBERT khi có GPU; thêm chấm điểm câu trả lời bằng người đọc."])

# ------------------------------------------------------------------ 8
H("8. Hướng dẫn chạy")
P("Windows: bấm đúp CHAY.bat. Hoặc:\npython -m venv .venv\n.venv\\Scripts\\activate     (Linux/macOS: source .venv/bin/activate)\npip install -r requirements.txt\npython app.py          # mở http://127.0.0.1:5000 (lần đầu 1–3 phút, đừng bấm Ctrl+C)\n"
  "python train.py        # (tùy chọn) huấn luyện lại + sinh biểu đồ\npython build_cache.py  # dự đoán sẵn cho 5.500 bài\npython predict.py \"nội dung bài báo\"\npython make_report.py  # sinh lại báo cáo\n"
  "python experiments/exp2_ablation.py   # chạy lại từng thí nghiệm (exp1 … exp8; cần pip install pyvi cho phần tách từ)")
P("Nếu gặp lỗi khi chạy, chạy python kiem_tra_moi_truong.py để xem phiên bản Python, thời gian nạp từng thư viện và thử nạp mô hình.")

# ------------------------------------------------------------------ 9
H("9. Kết luận")
P(f"Hệ thống TF-IDF + Linear SVM đạt Accuracy {pc(m['test_accuracy'])} và Macro-F1 {pc(m['test_macro_f1'])} trên 11 chủ đề (khoảng tin cậy 95% của Accuracy: {pc(ci['acc_lo'])}–{pc(ci['acc_hi'])}), "
  f"ổn định khi chia lại dữ liệu ({vn(sv_.split_acc_mean * 100, 1)}% ± {vn(sv_.split_acc_std * 100, 1)}), huấn luyện khoảng 10 giây trên CPU. "
  "Khi chỉnh tham số công bằng cho mọi mô hình, Linear SVM và Logistic Regression ngang nhau về mặt thống kê; ablation cho thấy nội dung bài báo, cặp từ và log-tf là các yếu tố quan trọng còn tách từ không cải thiện đáng kể. "
  f"Chúng em còn hiệu chỉnh độ tin cậy (ECE từ {pc(c0.ECE)} xuống {pc(c1.ECE)}), kiểm chứng rằng lời giải thích của AI Lab phản ánh đúng quyết định của mô hình trên 825 bài, "
  "và bổ sung chức năng Hỏi AI trả lời có dẫn nguồn. Toàn bộ chạy trên CPU, không cần tài nguyên tính toán lớn.")
H("Tài liệu tham khảo")
B(["Pedregosa et al., Scikit-learn: Machine Learning in Python, JMLR 12, 2011.", "Manning, Raghavan, Schütze. Introduction to Information Retrieval, Cambridge UP, 2008.",
   "Joachims, T. Text categorization with Support Vector Machines, ECML 1998.",
   "Dietterich, T. G. Approximate Statistical Tests for Comparing Supervised Classification Learning Algorithms, Neural Computation 10(7), 1998.",
   "Guo, C., Pleiss, G., Sun, Y., Weinberger, K. Q. On Calibration of Modern Neural Networks, ICML 2017.",
   "DeYoung, J. et al. ERASER: A Benchmark to Evaluate Rationalized NLP Models, ACL 2020.",
   "Mihalcea, R., Tarau, P. TextRank: Bringing Order into Texts, EMNLP 2004."])
d.save(Path(__file__).parent / "BAO_CAO_DO_AN.docx")
print("Đã sinh BAO_CAO_DO_AN.docx")
